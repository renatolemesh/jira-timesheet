"""Backend do dashboard de apontamentos do Jira.

Sobe um servidor FastAPI que expoe a API de relatorios e serve o front-end
estatico. O login e feito na Atlassian (OAuth 2.0 3LO, ver oauth.py): o
navegador so guarda um cookie de sessao, e o token fica em memoria aqui.
Credenciais do .env (API token) continuam como fallback para uso local.

Rodar:  python app.py   (ou: uvicorn app:app --reload)
"""

import csv
import hashlib
import html
import io
import mimetypes
import secrets
import os
import threading
import time
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from urllib.parse import urlencode

from fastapi import Cookie, Depends, FastAPI, HTTPException, Query, Request
from fastapi.responses import FileResponse, HTMLResponse, RedirectResponse, Response
from fastapi.staticfiles import StaticFiles

import oauth
from jira_client import JiraClient, JiraError

BASE_DIR = Path(__file__).parent

# a imagem python:slim nao tem /etc/mime.types: sem isto as capturas .webp da
# landing saem como application/octet-stream
mimetypes.add_type("image/webp", ".webp")


def load_env():
    # .env.oauth: client id/secret do app na Atlassian (no container chega via env_file)
    for name in (".env", ".env.oauth"):
        env_path = BASE_DIR / name
        if not env_path.exists():
            continue
        for line in env_path.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, value = line.split("=", 1)
            os.environ.setdefault(key.strip(), value.strip().strip('"').strip("'"))


load_env()

# fallback opcional: se o .env tiver credenciais, o app funciona sem login
ENV_BASE_URL = os.environ.get("JIRA_BASE_URL", "")
ENV_EMAIL = os.environ.get("JIRA_EMAIL", "")
ENV_TOKEN = os.environ.get("JIRA_TOKEN", "")
PORT = int(os.environ.get("PORT", "8000"))

app = FastAPI(title="Apontamentos de Horas")

# O callback OAuth esta cadastrado so no dominio rlhtech; o cookie de sessao
# criado la nao vale em outro host. O dominio antigo (o denunciado como
# phishing) so redireciona.
DOMINIOS_ANTIGOS = {"jira-timesheet.pharmaprices.shop"}
DOMINIO_CANONICO = "jira-timesheet.rlhtech.com.br"


@app.middleware("http")
async def redirecionar_dominio_antigo(request: Request, call_next):
    if request.headers.get("host", "").split(":")[0] in DOMINIOS_ANTIGOS:
        destino = f"https://{DOMINIO_CANONICO}{request.url.path}"
        if request.url.query:
            destino += f"?{request.url.query}"
        return RedirectResponse(destino, status_code=301)
    return await call_next(request)


# ---------- credenciais por requisicao ----------

_env_client = None
_env_client_lock = threading.Lock()


def env_client():
    """Fallback para uso local: credenciais fixas no .env (API token)."""
    global _env_client
    if not (ENV_BASE_URL and ENV_EMAIL and ENV_TOKEN):
        return None
    base_url = ENV_BASE_URL.strip().rstrip("/")
    with _env_client_lock:
        if _env_client is None:
            _env_client = JiraClient(base_url, ENV_EMAIL.strip(), ENV_TOKEN.strip())
    fp = hashlib.sha256(f"{base_url}|{ENV_EMAIL}".encode()).hexdigest()[:16]
    return _env_client, fp, base_url


def client_dep(sessao: str | None = Cookie(None)):
    """(client, fingerprint, url do site) da sessao OAuth - ou do .env."""
    s = oauth.obter_sessao(sessao)
    if s is not None and s.site is not None:
        try:
            s.garantir_token()
        except oauth.OAuthError:
            # refresh_token expirado ou revogado em id.atlassian.com
            oauth.encerrar_sessao(sessao)
            raise HTTPException(401, "Sessao expirada — entre novamente.")
        return s.client, s.fingerprint, s.site["url"]
    ctx = env_client()
    if ctx is None:
        raise HTTPException(401, "Entre com a sua conta Atlassian.")
    return ctx


def jira_http_error(e: JiraError) -> HTTPException:
    msg = str(e)
    if "-> 401" in msg or "-> 403" in msg:
        return HTTPException(401, "Jira recusou as credenciais — faca login novamente.")
    return HTTPException(502, msg)

# ---------- cache simples com TTL ----------

_cache: dict = {}
_cache_lock = threading.Lock()
CACHE_TTL = 300  # 5 min


def cached(key, builder, refresh: bool = False, ttl: int = CACHE_TTL):
    now = time.time()
    with _cache_lock:
        hit = _cache.get(key)
        if hit and not refresh and now - hit[0] < ttl:
            return hit[1]
    data = builder()
    with _cache_lock:
        _cache[key] = (time.time(), data)
    return data


# ---------- helpers ----------


def parse_date(value: str, name: str) -> date:
    try:
        return date.fromisoformat(value)
    except (TypeError, ValueError):
        raise HTTPException(400, f"Parametro '{name}' invalido: esperado YYYY-MM-DD")


def csv_param(value: str | None) -> list[str]:
    if not value:
        return []
    return [v.strip() for v in value.split(",") if v.strip()]


def day_ms(d: date) -> int:
    return int(datetime(d.year, d.month, d.day, tzinfo=timezone.utc).timestamp() * 1000)


def adf_text(node) -> str:
    """Extrai texto puro de um documento ADF (comentario de worklog)."""
    if not node:
        return ""
    if isinstance(node, str):
        return node
    parts = []
    if isinstance(node, dict):
        if node.get("type") == "text":
            parts.append(node.get("text", ""))
        for child in node.get("content") or []:
            child_text = adf_text(child)
            if child_text:
                parts.append(child_text)
    return " ".join(parts)


def quote_list(values: list[str]) -> str:
    return ", ".join('"' + v.replace('"', "") + '"' for v in values)


def business_days(start: date, end: date) -> list[date]:
    days = []
    d = start
    while d <= end:
        if d.weekday() < 5:  # seg a sex
            days.append(d)
        d += timedelta(days=1)
    return days


# ---------- montagem do relatorio ----------

ISSUE_FIELDS = ["summary", "project", "status", "issuetype", "assignee", "timetracking"]


def issue_info(issue: dict) -> dict:
    f = issue.get("fields", {})
    status = f.get("status") or {}
    assignee = f.get("assignee") or {}
    tt = f.get("timetracking") or {}
    return {
        "issueId": str(issue.get("id")),
        "issueKey": issue.get("key"),
        "summary": f.get("summary") or "",
        "projectKey": (f.get("project") or {}).get("key") or "",
        "projectName": (f.get("project") or {}).get("name") or "",
        "status": status.get("name") or "",
        "statusCategory": (status.get("statusCategory") or {}).get("key") or "",
        "issueType": (f.get("issuetype") or {}).get("name") or "",
        "assigneeId": assignee.get("accountId") or "",
        "assigneeName": assignee.get("displayName") or "",
        "estimateSeconds": tt.get("originalEstimateSeconds") or 0,
        "totalSpentSeconds": tt.get("timeSpentSeconds") or 0,
    }


def fetch_report(
    client: JiraClient, start: date, end: date, users: list[str], projects: list[str]
) -> dict:
    clauses = [f'worklogDate >= "{start.isoformat()}" AND worklogDate <= "{end.isoformat()}"']
    if projects:
        clauses.append(f"project in ({quote_list(projects)})")
    if users:
        clauses.append(f"worklogAuthor in ({quote_list(users)})")
    jql = " AND ".join(clauses)

    issues = client.search_issues(jql, ISSUE_FIELDS)
    info_by_id = {str(i["id"]): issue_info(i) for i in issues}

    # margem de +-1 dia para nao perder worklogs por fuso horario;
    # o filtro exato e feito abaixo pela data local do 'started'
    after_ms = day_ms(start - timedelta(days=1))
    before_ms = day_ms(end + timedelta(days=2))
    logs_by_issue = client.worklogs_for_issues(list(info_by_id), after_ms, before_ms)

    start_s, end_s = start.isoformat(), end.isoformat()
    user_set = set(users)
    entries = []
    for iid, logs in logs_by_issue.items():
        info = info_by_id.get(iid, {})
        for wl in logs:
            started = wl.get("started") or ""
            wl_date = started[:10]
            if not (start_s <= wl_date <= end_s):
                continue
            author = wl.get("author") or {}
            author_id = author.get("accountId") or ""
            if user_set and author_id not in user_set:
                continue
            entries.append(
                {
                    "date": wl_date,
                    "started": started,
                    "authorId": author_id,
                    "authorName": author.get("displayName") or "",
                    "seconds": wl.get("timeSpentSeconds") or 0,
                    "timeSpent": wl.get("timeSpent") or "",
                    "comment": adf_text(wl.get("comment")).strip(),
                    **{
                        k: info.get(k, "")
                        for k in (
                            "issueId",
                            "issueKey",
                            "summary",
                            "projectKey",
                            "projectName",
                            "status",
                            "statusCategory",
                            "issueType",
                            "estimateSeconds",
                            "totalSpentSeconds",
                        )
                    },
                }
            )

    entries.sort(key=lambda e: (e["date"], e["started"]))
    return {
        "start": start_s,
        "end": end_s,
        "entries": entries,
        "generatedAt": datetime.now(timezone.utc).isoformat(),
    }


def get_report(
    client: JiraClient, fp: str, start: date, end: date,
    users: list[str], projects: list[str], refresh: bool,
):
    key = ("report", fp, start.isoformat(), end.isoformat(),
           tuple(sorted(users)), tuple(sorted(projects)))
    return cached(key, lambda: fetch_report(client, start, end, users, projects), refresh)


# ---------- rotas ----------


# ---------- login com a Atlassian (OAuth 2.0 3LO) ----------


def _voltar_com_erro(msg: str) -> RedirectResponse:
    return RedirectResponse("/app?" + urlencode({"login_error": msg}), status_code=303)


def _escolher_site(s: "oauth.Sessao", site: dict) -> str | None:
    """Liga a sessao a um site Jira. Devolve mensagem de erro ou None."""
    client = JiraClient(oauth.API_BASE.format(cloud_id=site["id"]), bearer=s.access_token)
    try:
        me = client.myself()
    except JiraError:
        return f"Nao consegui ler o Jira de {site['name']} com a autorizacao concedida."
    s.site, s.client = site, client
    s.account_id = me.get("accountId", "")
    s.display_name = me.get("displayName", "")
    return None


def _cookie_sessao(resp: Response, sid: str):
    resp.set_cookie(
        oauth.COOKIE_SESSAO, sid, max_age=oauth.SESSAO_MAX_IDADE,
        httponly=True, secure=True, samesite="lax",
    )


@app.get("/auth/login")
def auth_login():
    if not oauth.configurado():
        return _voltar_com_erro("Login com a Atlassian nao configurado no servidor.")
    state = secrets.token_urlsafe(24)
    resp = RedirectResponse(oauth.url_autorizacao(state), status_code=303)
    # samesite=lax: o cookie precisa voltar no redirect da Atlassian para o callback
    resp.set_cookie(
        oauth.COOKIE_STATE, state, max_age=600, path="/auth",
        httponly=True, secure=True, samesite="lax",
    )
    return resp


@app.get("/auth/callback")
def auth_callback(
    code: str | None = None,
    state: str | None = None,
    error: str | None = None,
    oauth_state: str | None = Cookie(None),
):
    if error:
        return _voltar_com_erro("Autorizacao cancelada na Atlassian.")
    if not (code and state and oauth_state and secrets.compare_digest(state, oauth_state)):
        return _voltar_com_erro("O login expirou ou veio de outra aba — tente de novo.")
    try:
        tokens = oauth.trocar_code(code)
        sites = oauth.sites_autorizados(tokens["access_token"])
    except (oauth.OAuthError, KeyError):
        return _voltar_com_erro("A Atlassian recusou o login — tente de novo.")
    if not sites:
        return _voltar_com_erro("Nenhum site Jira foi autorizado para o app.")

    sid, s = oauth.criar_sessao(tokens, sites)
    destino = "/auth/sites"
    if len(sites) == 1:
        erro = _escolher_site(s, sites[0])
        if erro:
            oauth.encerrar_sessao(sid)
            return _voltar_com_erro(erro)
        destino = "/app"

    resp = RedirectResponse(destino, status_code=303)
    _cookie_sessao(resp, sid)
    resp.delete_cookie(oauth.COOKIE_STATE, path="/auth")
    return resp


@app.get("/auth/sites")
def auth_sites(sessao: str | None = Cookie(None)):
    """Escolha do site quando o usuario autorizou mais de um."""
    s = oauth.obter_sessao(sessao)
    if s is None:
        return RedirectResponse("/", status_code=303)
    botoes = "\n".join(
        f'<form method="post" action="/auth/site/{html.escape(site["id"])}">'
        f'<button class="btn primary" type="submit">{html.escape(site["name"])}</button> '
        f'<span class="muted">{html.escape(site["url"])}</span></form>'
        for site in s.sites
    )
    return HTMLResponse(f"""<!doctype html>
<html lang="pt-BR"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Escolha o site — Apontamentos de Horas</title>
<link rel="stylesheet" href="/static/style.css"><link rel="stylesheet" href="/static/legal.css">
</head><body><main class="legal">
<h1>Escolha o site do Jira</h1>
<p>Voce autorizou mais de um site. Qual deles quer consultar?</p>
<div style="display:grid;gap:12px">{botoes}</div>
</main></body></html>""")


@app.post("/auth/site/{cloud_id}")
def auth_site(cloud_id: str, sessao: str | None = Cookie(None)):
    s = oauth.obter_sessao(sessao)
    site = next((x for x in (s.sites if s else []) if x["id"] == cloud_id), None)
    if site is None:
        return RedirectResponse("/", status_code=303)
    erro = _escolher_site(s, site)
    if erro:
        return _voltar_com_erro(erro)
    return RedirectResponse("/app", status_code=303)


@app.post("/auth/logout")
def auth_logout(sessao: str | None = Cookie(None)):
    oauth.encerrar_sessao(sessao)
    resp = Response(status_code=204)
    resp.delete_cookie(oauth.COOKIE_SESSAO)
    return resp


@app.get("/api/meta")
def api_meta(refresh: bool = False, ctx=Depends(client_dep)):
    client, fp, base_url = ctx

    def build():
        users = client.users()
        projects = client.projects()
        me = client.myself()
        return {
            "baseUrl": base_url,
            "myself": {"accountId": me.get("accountId"), "displayName": me.get("displayName")},
            "users": sorted(
                (
                    {
                        "accountId": u["accountId"],
                        "displayName": u.get("displayName", ""),
                        "avatarUrl": (u.get("avatarUrls") or {}).get("24x24", ""),
                    }
                    for u in users
                ),
                key=lambda u: u["displayName"].lower(),
            ),
            "projects": sorted(
                ({"key": p["key"], "name": p.get("name", "")} for p in projects),
                key=lambda p: p["name"].lower(),
            ),
        }

    try:
        return cached(("meta", fp), build, refresh, ttl=3600)
    except JiraError as e:
        raise jira_http_error(e)


@app.get("/api/report")
def api_report(
    start: str = Query(...),
    end: str = Query(...),
    users: str | None = None,
    projects: str | None = None,
    refresh: bool = False,
    ctx=Depends(client_dep),
):
    client, fp, _ = ctx
    start_d, end_d = parse_date(start, "start"), parse_date(end, "end")
    if end_d < start_d:
        raise HTTPException(400, "'end' anterior a 'start'")
    if (end_d - start_d).days > 400:
        raise HTTPException(400, "Periodo maximo: 400 dias")
    try:
        return get_report(client, fp, start_d, end_d, csv_param(users), csv_param(projects), refresh)
    except JiraError as e:
        raise jira_http_error(e)


@app.get("/api/missing")
def api_missing(
    start: str = Query(...),
    end: str = Query(...),
    users: str | None = None,
    projects: str | None = None,
    expected: float = 8.0,
    refresh: bool = False,
    ctx=Depends(client_dep),
):
    """Dias uteis com apontamento abaixo do esperado + tasks mexidas sem apontamento."""
    client, fp, _ = ctx
    start_d, end_d = parse_date(start, "start"), parse_date(end, "end")
    user_ids = csv_param(users)
    project_keys = csv_param(projects)
    today = date.today()
    horizon = min(end_d, today)

    try:
        report = get_report(client, fp, start_d, end_d, user_ids, project_keys, refresh)

        # issues atualizadas no periodo (candidatas a "trabalhou mas nao apontou")
        clauses = [
            f'updated >= "{start_d.isoformat()}" AND updated < "{(end_d + timedelta(days=1)).isoformat()}"'
        ]
        if project_keys:
            clauses.append(f"project in ({quote_list(project_keys)})")
        clauses.append(
            f"assignee in ({quote_list(user_ids)})" if user_ids else "assignee is not EMPTY"
        )
        updated_issues = client.search_issues(" AND ".join(clauses), ISSUE_FIELDS)

        with _cache_lock:
            meta_hit = _cache.get(("meta", fp))
        meta = meta_hit[1] if meta_hit else {}
        name_by_id = {u["accountId"]: u["displayName"] for u in (meta.get("users") or [])}
    except JiraError as e:
        raise jira_http_error(e)

    # segundos apontados por (issue, autor) e por (autor, dia)
    by_issue_author: dict[tuple[str, str], int] = {}
    by_author_day: dict[str, dict[str, int]] = {}
    for e in report["entries"]:
        by_issue_author[(e["issueId"], e["authorId"])] = (
            by_issue_author.get((e["issueId"], e["authorId"]), 0) + e["seconds"]
        )
        by_author_day.setdefault(e["authorId"], {})
        by_author_day[e["authorId"]][e["date"]] = (
            by_author_day[e["authorId"]].get(e["date"], 0) + e["seconds"]
        )
        name_by_id.setdefault(e["authorId"], e["authorName"])

    # usuarios avaliados: filtro selecionado, senao autores + responsaveis vistos
    target_ids = list(user_ids)
    if not target_ids:
        seen = set(by_author_day)
        for i in updated_issues:
            assignee = (i.get("fields", {}).get("assignee") or {})
            if assignee.get("accountId"):
                seen.add(assignee["accountId"])
                name_by_id.setdefault(assignee["accountId"], assignee.get("displayName", ""))
        target_ids = sorted(seen, key=lambda a: (name_by_id.get(a) or "").lower())

    expected_seconds = int(expected * 3600)
    users_out = []
    workdays = business_days(start_d, horizon) if horizon >= start_d else []
    for uid in target_ids:
        days_logged = by_author_day.get(uid, {})
        missing_days = []
        for d in workdays:
            secs = days_logged.get(d.isoformat(), 0)
            if secs < expected_seconds:
                missing_days.append(
                    {"date": d.isoformat(), "seconds": secs, "isToday": d == today}
                )
        users_out.append(
            {
                "accountId": uid,
                "displayName": name_by_id.get(uid, uid),
                "totalSeconds": sum(days_logged.values()),
                "daysEvaluated": len(workdays),
                "missingDays": missing_days,
            }
        )

    issues_out = []
    seen_keys = set()
    for issue in updated_issues:
        info = issue_info(issue)
        if not info["assigneeId"] or info["assigneeId"] not in set(target_ids):
            continue
        if info["issueKey"] in seen_keys:
            continue
        seen_keys.add(info["issueKey"])
        own = by_issue_author.get((info["issueId"], info["assigneeId"]), 0)
        if own > 0:
            continue
        others = sum(
            secs for (iid, aid), secs in by_issue_author.items()
            if iid == info["issueId"] and aid != info["assigneeId"]
        )
        issues_out.append({**info, "othersSeconds": others})

    category_order = {"indeterminate": 0, "new": 1, "done": 2}
    issues_out.sort(
        key=lambda i: (category_order.get(i["statusCategory"], 3), i["assigneeName"].lower())
    )

    return {
        "start": start_d.isoformat(),
        "end": end_d.isoformat(),
        "horizon": horizon.isoformat(),
        "expectedHours": expected,
        "users": users_out,
        "issuesWithoutWorklog": issues_out,
    }


# ---------- exportacao (CSV / Excel) ----------


def fmt_br(iso: str) -> str:
    y, m, d = iso.split("-")
    return f"{d}/{m}/{y}"


def fmt_hm(seconds: int) -> str:
    h, m = seconds // 3600, (seconds % 3600) // 60
    return f"{h}h{m:02d}" if m else f"{h}h"


def fmt_brl(value: float) -> str:
    s = f"{value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"R$ {s}"


# (chave, rotulo, extrator) — a ordem aqui define a ordem das colunas
EXPORT_FIELDS = [
    ("data", "Data", lambda e: fmt_br(e["date"])),
    ("hora", "Hora", lambda e: (e.get("started") or "")[11:16]),
    ("projeto", "Projeto", lambda e: e["projectName"]),
    ("task", "Task", lambda e: e["issueKey"]),
    ("resumo", "Resumo", lambda e: e["summary"]),
    ("tipo", "Tipo", lambda e: e["issueType"]),
    ("status", "Status", lambda e: e["status"]),
    ("pessoa", "Pessoa", lambda e: e["authorName"]),
    ("tempo", "Tempo", lambda e: e["timeSpent"]),
    ("horas", "Horas", lambda e: round(e["seconds"] / 3600, 2)),
    ("comentario", "Comentário", lambda e: e["comment"]),
]
DEFAULT_EXPORT_FIELDS = ["data", "projeto", "task", "resumo", "pessoa", "tempo", "horas", "comentario"]


def build_csv_bytes(title, info, cols, rows, total_row, header_block) -> bytes:
    out = io.StringIO()
    w = csv.writer(out, delimiter=";", lineterminator="\r\n")

    def conv(v):
        return f"{v:.2f}".replace(".", ",") if isinstance(v, float) else v

    if header_block:
        w.writerow([title])
        for label, value in info:
            w.writerow([label, value])
        w.writerow([])
    w.writerow(cols)
    for row in rows:
        w.writerow([conv(v) for v in row])
    if total_row:
        w.writerow([conv(v) for v in total_row])
    return ("\ufeff" + out.getvalue()).encode("utf-8")  # BOM p/ Excel abrir como UTF-8


def build_xlsx_bytes(title, info, cols, rows, total_row, header_block, money_cols=()) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Relatório"
    r = 1
    if header_block:
        ws.cell(r, 1, title).font = Font(bold=True, size=13)
        r += 2
        for label, value in info:
            ws.cell(r, 1, label).font = Font(bold=True)
            ws.cell(r, 2, value)
            r += 1
        r += 1
    header_row = r
    for c, name in enumerate(cols, 1):
        ws.cell(r, c, name).font = Font(bold=True)
    r += 1
    money_fmt = '"R$" #,##0.00'
    for row in rows:
        for c, v in enumerate(row, 1):
            cell = ws.cell(r, c, v)
            if isinstance(v, float):
                cell.number_format = money_fmt if c in money_cols else "0.00"
        r += 1
    if total_row:
        for c, v in enumerate(total_row, 1):
            cell = ws.cell(r, c, v)
            cell.font = Font(bold=True)
            if isinstance(v, float):
                cell.number_format = money_fmt if c in money_cols else "0.00"
    ws.freeze_panes = ws.cell(header_row + 1, 1)

    # largura das colunas pelo conteudo (com teto)
    for c in range(1, len(cols) + 1):
        longest = len(str(cols[c - 1]))
        for row in rows[:200]:
            longest = max(longest, len(str(row[c - 1])))
        ws.column_dimensions[get_column_letter(c)].width = min(max(longest + 2, 10), 60)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@app.get("/api/export")
def api_export(
    start: str = Query(...),
    end: str = Query(...),
    users: str | None = None,
    projects: str | None = None,
    q: str | None = None,
    fields: str | None = None,
    fmt: str = "xlsx",
    header: bool = True,
    mode: str = "detalhado",
    rate: float | None = Query(None, ge=0),
    refresh: bool = False,
    ctx=Depends(client_dep),
):
    """Exporta o relatorio do periodo em CSV ou Excel, com cabecalho de totais.

    Se 'rate' (valor/hora) for informado, adiciona a coluna Valor (R$) e os
    totais em reais no cabecalho — pronto para emissao de fatura.
    """
    client, fp, _ = ctx
    start_d, end_d = parse_date(start, "start"), parse_date(end, "end")
    try:
        report = get_report(client, fp, start_d, end_d, csv_param(users), csv_param(projects), refresh)
    except JiraError as e:
        raise jira_http_error(e)

    entries = report["entries"]
    if q:
        needle = q.strip().lower()
        entries = [
            e for e in entries
            if needle in " ".join(
                [e["issueKey"], e["summary"], e["authorName"], e["comment"], e["projectName"], e["date"]]
            ).lower()
        ]

    total = sum(e["seconds"] for e in entries)
    people = sorted({e["authorName"] for e in entries if e["authorName"]})
    projs = sorted({e["projectName"] or e["projectKey"] for e in entries})
    rate = rate or 0

    def row_value(seconds: int) -> float:
        return round(seconds / 3600 * rate, 2)

    info = [
        ("Período", f"{fmt_br(report['start'])} a {fmt_br(report['end'])}"),
        ("Pessoas", ", ".join(people) if people else "—"),
        ("Projetos", ", ".join(projs) if projs else "—"),
        ("Apontamentos", str(len(entries))),
        ("Total de horas", f"{fmt_hm(total)} ({str(round(total / 3600, 2)).replace('.', ',')})"),
    ]

    if mode == "tasks":
        title = "Resumo por task — Jira"
        cols = ["Task", "Resumo", "Projeto", "Status", "Pessoas", "Apontamentos", "Horas", "Estimado (h)"]
        agg: dict[str, dict] = {}
        for e in entries:
            a = agg.setdefault(e["issueKey"], {
                "resumo": e["summary"], "projeto": e["projectName"], "status": e["status"],
                "pessoas": set(), "count": 0, "secs": 0, "estimado": e["estimateSeconds"],
            })
            a["count"] += 1
            a["secs"] += e["seconds"]
            a["pessoas"].add(e["authorName"])
        rows = [
            [key, a["resumo"], a["projeto"], a["status"], ", ".join(sorted(a["pessoas"])),
             a["count"], round(a["secs"] / 3600, 2),
             round(a["estimado"] / 3600, 2) if a["estimado"] else ""]
            for key, a in sorted(agg.items(), key=lambda kv: -kv[1]["secs"])
        ]
        total_row = ["TOTAL", "", "", "", "", len(entries), round(total / 3600, 2), ""]
        if rate:
            cols.append("Valor (R$)")
            ordered = sorted(agg.items(), key=lambda kv: -kv[1]["secs"])
            for row, (_, a) in zip(rows, ordered):
                row.append(row_value(a["secs"]))
            total_row.append(round(sum(r[-1] for r in rows), 2))
        base_name = "tasks"
    else:
        title = "Relatório de apontamentos — Jira"
        wanted = set(csv_param(fields) or DEFAULT_EXPORT_FIELDS)
        selected = [f for f in EXPORT_FIELDS if f[0] in wanted] or \
                   [f for f in EXPORT_FIELDS if f[0] in DEFAULT_EXPORT_FIELDS]
        cols = [label for _, label, _ in selected]
        keys = [key for key, _, _ in selected]
        rows = [[getter(e) for _, _, getter in selected] for e in entries]
        if rate:
            cols.append("Valor (R$)")
            for row, e in zip(rows, entries):
                row.append(row_value(e["seconds"]))
        total_row = [""] * len(cols)
        total_row[0] = "TOTAL"
        if "tempo" in keys:
            total_row[keys.index("tempo")] = fmt_hm(total)
        if "horas" in keys:
            total_row[keys.index("horas")] = round(total / 3600, 2)
        if rate:
            total_row[-1] = round(sum(r[-1] for r in rows), 2)
        base_name = "apontamentos"

    money_cols = ()
    if rate:
        total_value = total_row[-1] if rows else 0.0
        info.append(("Valor hora", fmt_brl(rate)))
        info.append(("Valor total", fmt_brl(total_value)))
        money_cols = (len(cols),)

    filename = f"{base_name}_{report['start']}_a_{report['end']}"
    if fmt == "xlsx":
        content = build_xlsx_bytes(title, info, cols, rows, total_row, header, money_cols)
        media = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        filename += ".xlsx"
    else:
        content = build_csv_bytes(title, info, cols, rows, total_row, header)
        media = "text/csv; charset=utf-8"
        filename += ".csv"
    return Response(
        content=content,
        media_type=media,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ---------- front-end ----------

app.mount("/static", StaticFiles(directory=BASE_DIR / "static"), name="static")


@app.get("/")
def landing(sessao: str | None = Cookie(None)):
    """Pagina publica; quem ja tem sessao vai direto para o painel."""
    s = oauth.obter_sessao(sessao)
    if s is not None and s.site is not None:
        return RedirectResponse("/app", status_code=303)
    return FileResponse(BASE_DIR / "static" / "landing.html")


@app.get("/app")
def index():
    return FileResponse(BASE_DIR / "static" / "index.html")


# paginas institucionais exigidas pelo cadastro do app OAuth na Atlassian
@app.get("/privacy")
def privacy():
    return FileResponse(BASE_DIR / "static" / "privacy.html")


@app.get("/terms")
def terms():
    return FileResponse(BASE_DIR / "static" / "terms.html")


@app.get("/contact")
def contact():
    return FileResponse(BASE_DIR / "static" / "contact.html")


if __name__ == "__main__":
    import uvicorn

    print(f"Dashboard em http://localhost:{PORT}")
    uvicorn.run(app, host="127.0.0.1", port=PORT)
